"""
Aelix backend  —  Flask serverless function for Vercel.

Serves the single-page app in aelix_login.html:
  1. POST /api/parse-pdf  -> extracts a heading/paragraph "sections" structure from a PDF
  2. POST /api/ai         -> edits a single section OR composes a new section from the whole doc
  3. POST /api/export     -> renders the (possibly AI-edited) sections back out as docx/pdf/html/md
"""

import io
import os
import statistics
from dotenv import load_dotenv
load_dotenv()

from flask import Flask, request, jsonify, send_file, send_from_directory
import pdfplumber
from groq import Groq
from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH

app = Flask(__name__, static_folder=".")

GROQ_MODEL_DEFAULT = "groq/compound-mini"
GROQ_FALLBACK_MODELS = ["groq/compound-mini", "groq/compound", "openai/gpt-oss-20b", "qwen/qwen3.6-27b"]


def call_groq_completion(client, preferred_model, messages, max_tokens):
    models_to_try = [preferred_model] + [m for m in GROQ_FALLBACK_MODELS if m != preferred_model]
    last_err = None
    for model in models_to_try:
        try:
            res = client.chat.completions.create(
                model=model,
                messages=messages,
                max_tokens=max_tokens
            )
            return res.choices[0].message.content.strip()
        except Exception as e:
            last_err = e
            print(f"Groq completion error with model {model}: {e}")
            continue
    raise last_err or Exception("All Groq models failed.")



@app.after_request
def add_cors(resp):
    resp.headers["Access-Control-Allow-Origin"] = "*"
    resp.headers["Access-Control-Allow-Headers"] = "Content-Type"
    resp.headers["Access-Control-Allow-Methods"] = "GET,POST,OPTIONS"
    return resp


# =========================================================================
#  PDF -> sections (heading/paragraph structure)
# =========================================================================
def pdf_to_sections(pdf_bytes):
    """
    Extract a heading/paragraph structure from a PDF using font-size heuristics.
    Returns a list of {"level": "h1"|"h2"|"h3"|"p", "content": str}, matching
    the shape the frontend (aelix_login.html) renders in its sidebar/editor.
    """
    sections = []
    try:
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            all_lines = []  # (text, size)
            for page in pdf.pages:
                words = page.extract_words(extra_attrs=["size"])
                if not words:
                    continue
                lines = {}
                for w in words:
                    key = round(w["top"], 1)
                    lines.setdefault(key, []).append(w)
                for top in sorted(lines.keys()):
                    line_words = sorted(lines[top], key=lambda w: w["x0"])
                    text = " ".join(w["text"] for w in line_words).strip()
                    if not text:
                        continue
                    size = max(w["size"] for w in line_words)
                    all_lines.append((text, size))

            if not all_lines:
                return []

            sizes = [s for _, s in all_lines]
            body_size = statistics.median(sizes)
            h1_t, h2_t, h3_t = body_size * 1.6, body_size * 1.3, body_size * 1.12

            def level_for(size):
                if size >= h1_t:
                    return "h1"
                if size >= h2_t:
                    return "h2"
                if size >= h3_t:
                    return "h3"
                return "p"

            buf_text = []
            for text, size in all_lines:
                lvl = level_for(size)
                if lvl != "p":
                    if buf_text:
                        sections.append({"level": "p", "content": " ".join(buf_text)})
                        buf_text = []
                    sections.append({"level": lvl, "content": text})
                else:
                    buf_text.append(text)
            if buf_text:
                sections.append({"level": "p", "content": " ".join(buf_text)})
    except Exception as e:
        print(f"Section extraction error: {e}")
    return sections


# =========================================================================
#  Routes
# =========================================================================
@app.route("/api/health", methods=["GET"])
def health():
    return jsonify({
        "ok": True,
        "groq_configured": bool(os.environ.get("GROQ_API_KEY"))
    })


@app.route("/api/parse-pdf", methods=["POST", "OPTIONS"])
def parse_pdf():
    """Upload PDF -> {"sections": [{level, content}, ...]}"""
    if request.method == "OPTIONS":
        return ("", 204)

    if "file" not in request.files:
        return jsonify({"error": "No file uploaded."}), 400

    file = request.files["file"]
    pdf_bytes = file.read()
    if not pdf_bytes:
        return jsonify({"error": "File is empty."}), 400

    try:
        sections = pdf_to_sections(pdf_bytes)
        if not sections:
            return jsonify({"error": "Could not extract any text from this PDF."}), 400
        return jsonify({"sections": sections})
    except Exception as e:
        return jsonify({"error": f"Error: {str(e)}"}), 400


@app.route("/api/ai", methods=["POST", "OPTIONS"])
def ai_command():
    """
    Two modes, matching what aelix_login.html sends:
      - Per-section edit:  {instruction, content, section} -> {"result": "<edited text>"}
      - Global compose:    {instruction, sections}         -> {"section": {level, content}}
    """
    if request.method == "OPTIONS":
        return ("", 204)

    data = request.get_json(silent=True) or {}
    instruction = (data.get("instruction") or "").strip()
    if not instruction:
        return jsonify({"error": "Missing instruction."}), 400

    api_key = os.environ.get("GROQ_API_KEY")
    if not api_key:
        return jsonify({"error": "GROQ_API_KEY not configured."}), 500

    client = Groq(api_key=api_key)
    model = os.environ.get("GROQ_MODEL", GROQ_MODEL_DEFAULT)

    # ---- Global action: compose one new section from the whole document ----
    if "sections" in data:
        sections = data.get("sections") or []
        context = "\n\n".join(
            f"[{(s.get('level') or 'p').upper()}] {s.get('content') or s.get('title') or ''}"
            for s in sections
        )[:8000]

        prompts = {
            "summarize": "Write a single concise paragraph (3-4 sentences) that summarises the document below.",
            "add-section": "Write one new paragraph that could naturally extend the document below.",
            "improve": "Write one paragraph that offers a refined, more polished restatement of the document's central point.",
            "tone": "Write one paragraph that restates the document's central point in a noticeably different tone.",
        }
        prompt = prompts.get(instruction, "Write one new paragraph related to the document below.")

        try:
            messages = [{
                "role": "user",
                "content": (
                    f"{prompt}\n\nReturn ONLY the paragraph text, no preamble, "
                    f"no markdown, no quotes.\n\nDocument:\n{context}"
                )
            }]
            result = call_groq_completion(client, model, messages, max_tokens=400)
            return jsonify({"section": {"level": "p", "content": result}})
        except Exception as e:
            return jsonify({"error": f"AI error: {str(e)}"}), 500

    # ---- Per-section edit ----
    content = (data.get("content") or "").strip()
    if not content:
        return jsonify({"error": "Missing content."}), 400
    level = data.get("section", "p")

    try:
        messages = [{
            "role": "user",
            "content": (
                f"Instruction: {instruction}\n\n"
                f"Edit the following {level} text according to the instruction. "
                f"Return ONLY the edited text, no preamble or explanation.\n\n"
                f"Text:\n{content}"
            )
        }]
        result = call_groq_completion(client, model, messages, max_tokens=1024)
        return jsonify({"result": result})
    except Exception as e:
        return jsonify({"error": f"AI error: {str(e)}"}), 500


@app.route("/api/export", methods=["POST", "OPTIONS"])
def export_file():
    """Export sections to Word, PDF, HTML or Markdown."""
    if request.method == "OPTIONS":
        return ("", 204)

    data = request.get_json(silent=True) or {}
    fmt = (data.get("format") or "docx").strip().lower()
    title = (data.get("title") or "Document").strip() or "Document"
    sections = data.get("sections", [])

    if not sections:
        return jsonify({"error": "No content to export."}), 400

    buf = io.BytesIO()

    # ========== DOCX (Word) ==========
    if fmt in ("docx", "doc", "word"):
        doc = DocxDocument()
        title_para = doc.add_heading(title, 0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        for s in sections:
            level = s.get("level", "p")
            text = (s.get("content") or "").strip()
            if not text:
                continue
            if level == "h1":
                doc.add_heading(text, level=1)
            elif level == "h2":
                doc.add_heading(text, level=2)
            elif level == "h3":
                doc.add_heading(text, level=3)
            else:
                doc.add_paragraph(text)

        doc.save(buf)
        buf.seek(0)
        return send_file(
            buf,
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            as_attachment=True,
            download_name=f"{title}.docx"
        )

    # ========== PDF ==========
    elif fmt == "pdf":
        from reportlab.lib.pagesizes import LETTER
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.units import inch

        doc = SimpleDocTemplate(buf, pagesize=LETTER)
        styles = getSampleStyleSheet()
        style_map = {"h1": "Heading1", "h2": "Heading2", "h3": "Heading3", "p": "Normal"}
        story = [Paragraph(title, styles["Title"]), Spacer(1, 0.3 * inch)]

        for s in sections:
            text = (s.get("content") or "").strip()
            if not text:
                continue
            style = styles[style_map.get(s.get("level", "p"), "Normal")]
            story.append(Paragraph(text, style))
            story.append(Spacer(1, 0.15 * inch))

        doc.build(story)
        buf.seek(0)
        return send_file(
            buf,
            mimetype="application/pdf",
            as_attachment=True,
            download_name=f"{title}.pdf"
        )

    # ========== HTML ==========
    elif fmt == "html":
        parts = [
            "<!DOCTYPE html><html><head><meta charset='utf-8'>"
            f"<title>{title}</title></head><body>",
            f"<h1>{title}</h1>"
        ]
        tag_map = {"h1": "h1", "h2": "h2", "h3": "h3"}
        for s in sections:
            text = (s.get("content") or "").strip()
            if not text:
                continue
            tag = tag_map.get(s.get("level", "p"), "p")
            parts.append(f"<{tag}>{text}</{tag}>")
        parts.append("</body></html>")

        buf.write("\n".join(parts).encode("utf-8"))
        buf.seek(0)
        return send_file(
            buf,
            mimetype="text/html",
            as_attachment=True,
            download_name=f"{title}.html"
        )

    # ========== Markdown ==========
    elif fmt == "md":
        prefix_map = {"h1": "# ", "h2": "## ", "h3": "### ", "p": ""}
        lines = [f"# {title}", ""]
        for s in sections:
            text = (s.get("content") or "").strip()
            if not text:
                continue
            lines.append(prefix_map.get(s.get("level", "p"), "") + text)
            lines.append("")

        buf.write("\n".join(lines).encode("utf-8"))
        buf.seek(0)
        return send_file(
            buf,
            mimetype="text/markdown",
            as_attachment=True,
            download_name=f"{title}.md"
        )

    return jsonify({"error": "Unsupported format."}), 400


# =========================================================================
#  Static & HTML Routes
# =========================================================================
@app.route("/", methods=["GET"])
def index():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    return send_file(os.path.join(base_dir, "aelix_login.html"))


@app.route("/editor", methods=["GET"])
def editor_page():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    return send_file(os.path.join(base_dir, "aelix_editor.html"))


@app.route("/<path:path>", methods=["GET"])
def static_proxy(path):
    base_dir = os.path.dirname(os.path.abspath(__file__))
    file_path = os.path.join(base_dir, path)
    if os.path.exists(file_path) and os.path.isfile(file_path):
        return send_file(file_path)
    return jsonify({"error": "Not found."}), 404


if __name__ == "__main__":
    app.run(port=5000, debug=True)

